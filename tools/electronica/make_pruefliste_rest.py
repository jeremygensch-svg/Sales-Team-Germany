"""Erzeugt electronica/pruefliste_rest.md: alle Zielfirmen (Prio A/B/C), die noch nicht im Browser geprüft sind."""
import os,re,sys
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
from openpyxl import load_workbook
from elparse import load,nk,ROOT
from aaaparse import demak,cut
os.chdir(ROOT)
wb=load_workbook('aussendienst/electronica_2026_Aussendienst.xlsx'); ws=wb['Zielfirmen']
todo=[r for r in ws.iter_rows(min_row=5,values_only=True) if r[0] in 'ABC' and str(r[7]).startswith('vorläufig')]
det={nk(x['firma'])[:10]:x for x in load()}
order={'A':0,'B':1,'C':2}; todo.sort(key=lambda r:(order[r[0]],str(r[3]),r[1].lower()))
L=["# electronica 2026 – Browser-Prüfliste für die restlichen Zielfirmen","",
f"Stand {__import__('datetime').date.today()}. {len(todo)} Firmen aus der Außendienst-Excel (Prio A/B/C), die noch nicht im Browser geprüft sind (Spalte Prüfstand = vorläufig). Reihenfolge: Prio A (Grün-Kandidaten, vor allem Anlagenbauer), dann B, dann C, je Gebiet.","",
"## Startanweisung für den lokalen Claude (Chrome-Erweiterung aktiv)","",
"Vorgehen und Regeln wie in `electronica/pruefliste_gruen.md` (Startanweisung dort lesen) und `electronica/ergebnisse_browser/gruen_01.md` als Formatvorlage.",
"- Entscheidungen von Jey: Verbände/Prüfdienstleister/Forschung raus. Firmen im CRM (gleicher oder sehr ähnlicher Name) raus. **Anlagenbauer bleiben drin**; raus nur, wenn die Firma selbst Verguss-/Dosieranlagen baut oder vertreibt (E1, mit URL und Zitat).",
"- Pro Firma im Browser: Impressum/Handelsregister (Sitz), Fertigung am Sitz, Verguss/Klebung/Dichtung am Produkt, Patente (Google Patents + Espacenet), Wettbewerber-Referenzen, Konzern/mehrere Werke, CRM-Abgleich (`python3 crm_abgleich.py`).",
"- Pro Durchgang höchstens 10 Firmen in dieser Reihenfolge. Ergebnis: `electronica/ergebnisse_browser/rest_NN.md`, Format wie `gruen_NN.md`: Überschrift `## N. Firma (Stand)`, danach Zeile `- Ampel neu: GRÜN|GELB|ROT … Tabellenzeile D: Z` und bei Gelb `- OFFENER AUFTRAG …`.",
"- Nach jedem Durchgang: committen, pushen. Die Excel baut sich per `python3 tools/electronica/build_el.py` neu (liest gruen_*.md und rest_*.md automatisch; Ampel, Prio und Spalte Prüfstand werden übernommen). Dann `python3 tools/electronica/make_pruefliste_rest.py` für die verbleibende Liste.",
"- Keine Kontaktsuche (Regel K), nichts verschicken.","","## Firmen","","| Nr | Prio | Firma | Stand | Gebiet |","|---|---|---|---|---|"]
for i,r in enumerate(todo,1): L.append(f"| {i} | {r[0]} | {r[1]} | {r[2]} | {r[3]} |")
L+=["","## Details je Firma",""]
for i,r in enumerate(todo,1):
    d=det.get(nk(r[1])[:10],{})
    nach=re.sub(r'^\(?\d\)?\s*','',d.get('nach','')) if d else ''
    L+=[f"### {i}. {r[1]} ({r[2]}) – Prio {r[0]}",
    f"- Gebiet/Ort: {r[3]} – {r[4]}"+(f" | Quellen: `electronica/ergebnisse/tom_batch_{d['b']:02d}.md`, `tina_batch_{d['b']:02d}.md`" if d else ''),
    f"- Was wird gebaut: {cut(str(r[5]),240)}",
    f"- Demak-Bereich: {r[6]}",
    f"- Offener Auftrag (Tina): {cut(nach,400) if nach else 'Fertigung vor Ort und Verguss/Klebung belegen; Impressum, Stellenanzeigen, Datenblätter lesen.'}",
    "- Zu prüfen: [ ] Impressum/Sitz  [ ] Fertigung am Sitz  [ ] Verguss/Klebung/Dichtung  [ ] Patente  [ ] Wettbewerbereinsatz  [ ] Konzern/Werke  [ ] CRM wiederholt  [ ] E1 (eigene Verguss-/Dosieranlagen?)",""]
open('electronica/pruefliste_rest.md','w').write('\n'.join(L)+'\n'); print(len(todo))
